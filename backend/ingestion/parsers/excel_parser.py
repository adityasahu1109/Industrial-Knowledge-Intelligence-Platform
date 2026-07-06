import openpyxl

def parse_excel(file_path: str) -> list[dict]:
    """
    Extracts text from an Excel file (.xlsx).
    Returns: list of dicts {"text": str, "page": int, "section": str}
    Here 'page' refers to the worksheet index.
    """
    pages = []
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    for sheet_idx, sheet_name in enumerate(wb.sheetnames):
        sheet = wb[sheet_name]
        current_text = []
        
        for row in sheet.iter_rows(values_only=True):
            # Convert row to string, ignoring None values
            row_vals = [str(val) for val in row if val is not None]
            if row_vals:
                current_text.append(" | ".join(row_vals))
                
        if current_text:
            pages.append({
                "text": "\n".join(current_text),
                "page": sheet_idx + 1,
                "section": sheet_name
            })
            
    return pages
